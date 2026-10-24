from openpyxl import load_workbook

from exampipe.export import HEADERS, write_bank
from exampipe.merge import build_bank
from exampipe.sample import clean_rows


def read(path):
    ws = load_workbook(path).active
    return [[c.value for c in r] for r in ws.iter_rows()]


def test_one_workbook_per_nonempty_subsection(course, tmp_path):
    bank = build_bank({"A": clean_rows()}, course)
    written = write_bank(bank, course, tmp_path)
    assert len(written) == 12
    assert (tmp_path / "Mathematics" / "01 Percentage.xlsx").exists()
    assert not (tmp_path / "_unmapped.xlsx").exists()


def test_empty_subsections_get_no_file(course, tmp_path):
    bank = build_bank({"A": clean_rows()[:3]}, course)
    assert len(write_bank(bank, course, tmp_path)) == 1


def test_contents_header_and_provenance(course, tmp_path):
    bank = build_bank({"P1": clean_rows()}, course)
    write_bank(bank, course, tmp_path)
    rows = read(tmp_path / "Mathematics" / "01 Percentage.xlsx")
    assert rows[0] == HEADERS
    assert len(rows) == 1 + 3
    first = rows[1]
    assert first[0] == "P1" and first[1] == 1 and first[3] == "Percentage"
    assert first[9] == "3"  # Correct Answer stays text


def test_uncertain_is_preserved_as_text(course, tmp_path, row):
    row.correct_answer = "UNCERTAIN"
    write_bank(build_bank({"A": [row]}, course), course, tmp_path)
    rows = read(tmp_path / "Mathematics" / "01 Percentage.xlsx")
    assert rows[1][9] == "UNCERTAIN"


def test_text_starting_with_equals_is_not_a_formula(course, tmp_path, row):
    row.question = "=1+1"
    write_bank(build_bank({"A": [row]}, course), course, tmp_path)
    ws = load_workbook(tmp_path / "Mathematics" / "01 Percentage.xlsx").active
    cell = ws["E2"]
    assert cell.value == "=1+1"
    assert cell.data_type == "s"


def test_illegal_control_characters_are_stripped(course, tmp_path, row):
    row.question = "bad\x00char\x07 here"
    write_bank(build_bank({"A": [row]}, course), course, tmp_path)
    rows = read(tmp_path / "Mathematics" / "01 Percentage.xlsx")
    assert rows[1][4] == "badchar here"


def test_unmapped_rows_go_to_their_own_file(course, tmp_path, row):
    row.sub_section = "Astrology"
    written = write_bank(build_bank({"A": [row]}, course), course, tmp_path)
    assert [p.name for p in written] == ["_unmapped.xlsx"]
    assert read(tmp_path / "_unmapped.xlsx")[1][3] == "Astrology"


def test_header_is_frozen(course, tmp_path):
    write_bank(build_bank({"A": clean_rows()}, course), course, tmp_path)
    ws = load_workbook(tmp_path / "Mathematics" / "01 Percentage.xlsx").active
    assert ws.freeze_panes == "A2"
