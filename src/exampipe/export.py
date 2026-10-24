"""Write a Bank to Excel files, one workbook per sub-section."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font

from .config import Course
from .merge import Bank, BankEntry

HEADERS = [
    "Paper", "No", "Section", "Sub-Section", "Question",
    "Answer 1", "Answer 2", "Answer 3", "Answer 4",
    "Correct Answer", "Explanation", "Difficulty",
]
_WIDTHS = [18, 6, 20, 24, 60, 22, 22, 22, 22, 14, 70, 12]
_WRAP_COLUMNS = {5, 11}  # Question, Explanation (1-based)

_UNSAFE_IN_FILENAME = re.compile(r'[\\/:*?"<>|]')


def _safe_name(text: str) -> str:
    return _UNSAFE_IN_FILENAME.sub("-", text).strip()


def _entry_values(entry: BankEntry) -> list:
    r = entry.row
    return [entry.paper, r.no, r.section, r.sub_section, r.question, *r.options,
            r.correct_answer, r.explanation, r.difficulty]


def _write_sheet(path: Path, entries: list[BankEntry]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(HEADERS)
    for col, (cell, width) in enumerate(zip(ws[1], _WIDTHS), start=1):
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"

    for entry in entries:
        values = _entry_values(entry)
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in values])
        for col, cell in enumerate(ws[ws.max_row], start=1):
            # Text that starts with "=" would be stored as a formula and
            # evaluated when the file is opened. Force it to plain text.
            if cell.data_type == "f":
                cell.data_type = "s"
            if col in _WRAP_COLUMNS:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_bank(bank: Bank, course: Course, out_dir: str | Path) -> list[Path]:
    """Write one workbook per non-empty sub-section, plus ``_unmapped.xlsx``
    if any rows could not be routed. Returns the files written."""
    out_dir = Path(out_dir)
    written: list[Path] = []
    for section in course.sections:
        for sub in section.subsections:
            entries = bank.entries.get((section.key, sub.serial))
            if not entries:
                continue
            path = out_dir / _safe_name(section.name) / f"{sub.serial:02d} {_safe_name(sub.name)}.xlsx"
            _write_sheet(path, entries)
            written.append(path)
    if bank.unmapped:
        path = out_dir / "_unmapped.xlsx"
        _write_sheet(path, bank.unmapped)
        written.append(path)
    return written
